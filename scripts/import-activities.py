"""Normalize Domingo activity workbooks into the runtime JSON fixture.

Usage:
  python scripts/import-activities.py <main.xlsx> <ideas.xlsx> [output.json]

The script intentionally keeps unknown factual values nullable. It requires
openpyxl only at import time; the application never reads Excel at runtime.
"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as error:  # pragma: no cover - import environment guard
    raise SystemExit("Install openpyxl to run the activity import") from error


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "src" / "data" / "fixtures" / "activities.generated.json"


def clean(value: Any) -> str | None:
    if value is None:
        return None
    text = unicodedata.normalize("NFKC", str(value)).strip()
    return text or None


def key(value: Any) -> str:
    return (clean(value) or "").casefold().replace("ё", "е")


def split_tags(value: Any) -> list[str]:
    parts = [key(part) for part in re.split(r"[;,]", clean(value) or "")]
    return list(dict.fromkeys(part for part in parts if part))


def matched_tags(value: Any, dictionary: dict[str, tuple[str, ...]]) -> list[str]:
    source = split_tags(value)
    result: list[str] = []
    for canonical, aliases in dictionary.items():
        if any(any(alias in part for alias in aliases) for part in source):
            result.append(canonical)
    return result


def minute_range(value: Any) -> dict[str, int] | None:
    if isinstance(value, (int, float)):
        minutes = max(0, int(value))
        return {"min": minutes, "max": minutes}
    numbers = [int(number) for number in re.findall(r"\d+", clean(value) or "")]
    if not numbers:
        return None
    return {"min": numbers[0], "max": numbers[-1]}


def integer(value: Any) -> int | None:
    parsed = minute_range(value)
    return parsed["min"] if parsed else None


SEASONS = {
    "spring": ("весн",),
    "summer": ("лет",),
    "autumn": ("осен",),
    "winter": ("зим",),
    "all": ("круглый год", "любое время года"),
}
WEATHER = {
    "sunny": ("солнеч",),
    "hot": ("жарк",),
    "cloudy": ("облач",),
    "cool": ("прохлад",),
    "rain": ("дожд",),
    "snow": ("снег",),
    "frost": ("мороз",),
    "dry": ("сух",),
    "warm": ("тепл",),
    "any": ("любая", "любую"),
}
MOODS = {
    "calm": ("спокой",),
    "active": ("актив",),
    "relax": ("расслаб", "восстанов"),
    "nature": ("природ", "прогул"),
    "food": ("вкус", "гастроном"),
    "learn": ("узнать", "познав", "истори", "культур"),
    "beautiful": ("красив", "вид", "фото"),
    "creative": ("твор", "мастер-класс"),
    "romantic": ("роман",),
    "special": ("особенн", "празд"),
}
COMPANIONS = {
    "solo": ("одному", "одна", "один"),
    "couple": ("вдвоем", "пара"),
    "children": ("с детьми", "семья", "дети"),
    "friends": ("друз", "компан"),
    "teens": ("подрост",),
    "pet": ("питом", "с собак"),
}
TIMES = {
    "morning": ("утро",),
    "day": ("день",),
    "evening": ("вечер",),
    "any": ("любое",),
}


def booking_requirement(value: Any) -> str:
    text = key(value)
    if not text or text == "нет":
        return "none"
    if "уточ" in text:
        return "check"
    if "желатель" in text or "рекомен" in text:
        return "recommended"
    if "да" in text or "обяз" in text or "заранее" in text:
        return "required"
    return "check"


def location_group(value: Any) -> str:
    text = key(value)
    groups = (
        ("Domingo", ("domingo", "дом ", "территор", "венский лес", "нара вилладж", "игнатьево", "берендеево")),
        ("Серпухов", ("серпухов",)),
        ("Таруса", ("тарус",)),
        ("Поленово", ("поленов",)),
        ("Мелихово", ("мелихов",)),
        ("Чехов", ("чехов",)),
    )
    for group, aliases in groups:
        if any(alias in text for alias in aliases):
            return group
    return clean(value) or "Не указано"


def activity_type(section: Any, title: Any, source: Any) -> str:
    text = " ".join((key(section), key(title), key(source)))
    if "маршрут" in text or "тропа" in text:
        return "route"
    if any(word in text for word in ("доставка", "массаж", "баня", "фурако", "завтрак", "кейтеринг")):
        return "service"
    if any(word in text for word in ("музей", "усадьб", "ресторан", "ферм", "spa")):
        return "place"
    return "activity"


def paid_service(title: Any, section: Any, source: Any) -> bool:
    text = " ".join((key(title), key(section)))
    return "domingo" in key(source) and any(
        word in text
        for word in ("фурако", "баня", "массаж", "завтрак", "доставка", "sup", "питание", "кейтеринг", "торт", "цвет")
    )


def normalize_row(row: dict[str, Any], dataset: str) -> dict[str, Any]:
    if dataset == "main":
        source = row.get("Источник")
        title = row.get("Активность / место")
        description = row.get("Короткое описание")
        duration = row.get("Длительность на месте, мин")
        total = row.get("Общий бюджет времени, мин")
        mood = row.get("Тип отдыха / настроение")
        map_asset = row.get("Карта / маршрут / скрин")
        source_url = row.get("Ссылка-источник")
        enabled = key(row.get("Включать в подбор"))
        status = "active" if enabled == "да" else "disabled" if enabled == "нет" else "draft"
        priority = None
    else:
        source = row.get("Тип")
        title = row.get("Место / идея")
        description = row.get("Описание")
        duration = row.get("Длительность, мин")
        total = row.get("Общий бюджет времени, мин")
        mood = row.get("Настроение / тип отдыха")
        map_asset = row.get("Карта / материал")
        source_url = row.get("Ссылка")
        status = "active" if key(row.get("Добавить в основную базу")) == "да" else "draft"
        priority = clean(row.get("Приоритет"))

    section = row.get("Раздел")
    location = row.get("Локация")
    booking = row.get("Запись заранее")
    item_id = clean(row.get("ID"))
    if not item_id or not clean(title):
        raise ValueError("Every imported row must have ID and title")

    seasons = matched_tags(row.get("Сезон"), SEASONS)
    return {
        "id": item_id.lower(),
        "status": status,
        "source": "candidate" if dataset == "ideas" else ("domingo" if "domingo" in key(source) else "guide"),
        "category": clean(section) or "Без категории",
        "type": activity_type(section, title, source),
        "title": clean(title),
        "description": clean(description),
        "location": clean(location),
        "locationGroup": location_group(location),
        "travelMinutes": integer(row.get("Дорога в одну сторону, мин")),
        "durationMinutes": minute_range(duration),
        "totalMinutes": minute_range(total),
        "seasons": seasons or ["all"],
        "weather": matched_tags(row.get("Погода"), WEATHER) or ["any"],
        "conditions": clean(row.get("Температура / условия")),
        "timeOfDay": matched_tags(row.get("Время суток"), TIMES) or ["any"],
        "moods": matched_tags(mood, MOODS),
        "companions": matched_tags(row.get("Компания"), COMPANIONS),
        "ageRestrictions": clean(row.get("Возраст / ограничения")),
        "transport": split_tags(row.get("Транспорт")),
        "bookingRequirement": booking_requirement(booking),
        "notes": clean(row.get("Примечание") if dataset == "main" else row.get("Что проверить / примечание")),
        "mapAsset": clean(map_asset),
        "sourceUrl": clean(source_url),
        "priority": priority,
        "isSelectable": status == "active",
        "isPaidService": paid_service(title, section, source),
        "serviceId": item_id.lower() if paid_service(title, section, source) else None,
    }


def read_sheet(path: Path, sheet_name: str, dataset: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(value) or "" for value in next(rows)]
    result = []
    for values in rows:
        row = dict(zip(headers, values, strict=False))
        if any(value is not None for value in values):
            result.append(normalize_row(row, dataset))
    return result


def main() -> None:
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    main_path = Path(sys.argv[1]).resolve()
    ideas_path = Path(sys.argv[2]).resolve()
    output_path = Path(sys.argv[3]).resolve() if len(sys.argv) > 3 else DEFAULT_OUTPUT
    activities = read_sheet(main_path, "Активности", "main")
    activities.extend(read_sheet(ideas_path, "Новые места и идеи", "ideas"))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(activities, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    active = sum(item["status"] == "active" for item in activities)
    print(f"Imported {len(activities)} activities ({active} active) to {output_path}")


if __name__ == "__main__":
    main()
